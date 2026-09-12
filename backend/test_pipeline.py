import os
from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

TEST_DIR = Path(__file__).resolve().parent.parent / "test_sample_files"
TEST_DIR.mkdir(parents=True, exist_ok=True)

def create_multi_page_pdf():
    pdf_path = TEST_DIR / "real_os_multipage.pdf"
    doc = pymupdf.open()

    # Page 1
    page1 = doc.new_page(width=595, height=842)
    page1.insert_text((50, 70), "Operating Systems - Unit 1: Process Management", fontsize=16)
    page1.insert_text((50, 110), "A process is an instance of a computer program that is being executed by one or many threads.", fontsize=11)
    page1.insert_text((50, 140), "Key states: New, Ready, Running, Waiting, Terminated.", fontsize=11)

    # Page 2
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 70), "Chapter 3: Deadlock Prevention and Avoidance", fontsize=16)
    page2.insert_text((50, 110), "Four Coffman conditions: Mutual Exclusion, Hold and Wait, No Preemption, Circular Wait.", fontsize=11)
    page2.insert_text((50, 140), "Banker's algorithm ensures dynamic avoidance by verifying safe states.", fontsize=11)

    # Page 3
    page3 = doc.new_page(width=595, height=842)
    page3.insert_text((50, 70), "Chapter 4: Memory Virtualization", fontsize=16)
    page3.insert_text((50, 110), "Paging eliminates external fragmentation through fixed-size physical frames.", fontsize=11)

    doc.save(str(pdf_path))
    doc.close()
    print(f"Created 3-page PDF: {pdf_path}")
    return pdf_path

def create_multi_slide_pptx():
    pptx_path = TEST_DIR / "real_graph_slides.pptx"
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    # Slide 1
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p = txBox.text_frame.add_paragraph()
    p.text = "Lecture 6: Graph Traversal Algorithms"
    p.font.size = Pt(28)
    p.font.bold = True
    p2 = txBox.text_frame.add_paragraph()
    p2.text = "Overview of Breadth-First Search and Depth-First Search trade-offs."
    p2.font.size = Pt(16)

    # Slide 2
    slide2 = prs.slides.add_slide(blank_layout)
    txBox2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p_s2 = txBox2.text_frame.add_paragraph()
    p_s2.text = "Slide 2: Breadth-First Search (BFS)"
    p_s2.font.size = Pt(24)
    p_s2_body = txBox2.text_frame.add_paragraph()
    p_s2_body.text = "Explores level by level using FIFO queue. Time complexity O(V+E), Space O(V)."

    # Slide 3
    slide3 = prs.slides.add_slide(blank_layout)
    txBox3 = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    p_s3 = txBox3.text_frame.add_paragraph()
    p_s3.text = "Slide 3: Depth-First Search (DFS)"
    p_s3.font.size = Pt(24)
    p_s3_body = txBox3.text_frame.add_paragraph()
    p_s3_body.text = "Dives deep along each branch using recursion or LIFO stack. Space complexity O(h)."

    prs.save(str(pptx_path))
    print(f"Created 3-slide PPTX: {pptx_path}")
    return pptx_path

def create_markdown_file():
    md_path = TEST_DIR / "real_dbms_normalization.md"
    content = """# Relational Database Normalization

Normalization organizes tables to reduce data redundancy and eliminate update anomalies.

## First Normal Form (1NF)
Each table column must contain atomic (indivisible) values.
No repeating groups or arrays allowed in tuples.

## Second Normal Form (2NF)
Must be in 1NF and contain no partial functional dependencies.
Every non-prime attribute must depend fully on the entire candidate key.

## Third Normal Form (3NF) & BCNF
Must be in 2NF and eliminate transitive functional dependencies (X -> Y and Y -> Z).
Boyce-Codd Normal Form (BCNF) strictly requires every determinant X to be a superkey.
"""
    md_path.write_text(content.strip(), encoding="utf-8")
    print(f"Created structured Markdown: {md_path}")
    return md_path

def create_plain_text_file():
    txt_path = TEST_DIR / "real_syllabus_review.txt"
    content = """CS301 Course Syllabus - Exam Checklist
Section 1: Operating System Architectures (Monolithic vs Microkernel).
Section 2: Concurrency Primitives (Semaphores, Mutexes, Spinlocks).
Section 3: Deadlock Coffman Conditions and Banker Safe State Vector Analysis.
Section 4: Virtual Memory Management (Page Tables, TLB, Page Replacement).
"""
    txt_path.write_text(content.strip(), encoding="utf-8")
    print(f"Created plain text file: {txt_path}")
    return txt_path

def create_scanned_note_image():
    img_path = TEST_DIR / "real_handwritten_note.png"
    # Create an image that resembles a student handwritten notebook note with legible text
    img = Image.new("RGB", (650, 300), color=(254, 252, 245))
    draw = ImageDraw.Draw(img)

    # Draw subtle notebook ruling lines
    for y in range(40, 300, 35):
        draw.line([(20, y), (630, y)], fill=(220, 226, 235), width=1)
    # Red left margin line
    draw.line([(80, 10), (80, 290)], fill=(245, 180, 180), width=2)

    # Add text
    draw.text((95, 45), "Prof. Notes: Deadlock Exam Rule", fill=(20, 40, 120))
    draw.text((95, 80), "1. Prevention: statically break circular wait order.", fill=(15, 30, 90))
    draw.text((95, 115), "2. Avoidance: dynamic Banker safe state test.", fill=(15, 30, 90))
    draw.text((95, 150), "Important: Avoidance requires max claim known in advance.", fill=(15, 30, 90))

    img.save(str(img_path))
    print(f"Created scanned note image: {img_path}")
    return img_path

if __name__ == "__main__":
    create_multi_page_pdf()
    create_multi_slide_pptx()
    create_markdown_file()
    create_plain_text_file()
    create_scanned_note_image()
